"""DOCX text extraction using python-docx."""

from __future__ import annotations

import io
from docx import Document


def extract_text_from_docx(docx_bytes: bytes) -> tuple[str, int]:
    """Extract text and approximate page count from a DOCX file.

    Returns:
        (raw_text, estimated_page_count)
    """
    doc = Document(io.BytesIO(docx_bytes))

    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    raw_text = "\n\n".join(paragraphs)

    # Estimate pages: ~3000 characters per page for legal docs
    estimated_pages = max(1, len(raw_text) // 3000)

    return raw_text, estimated_pages
