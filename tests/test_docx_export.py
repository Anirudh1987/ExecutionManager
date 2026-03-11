"""Tests for Word document export (advisory, redline, compliance tracker)."""

import pytest
from io import BytesIO

from docx import Document

from src.models.contract import Clause, ClauseType, Contract
from src.models.deal import Deal
from src.models.review import (
    AIFinding,
    ClauseReview,
    DealContext,
    HumanAnnotation,
    HumanVerdict,
    Review,
    ReviewStage,
    RiskLevel,
)
from src.engine.docx_exporter import (
    export_advisory_docx,
    export_redline_docx,
    export_compliance_tracker_docx,
)


def _make_deal():
    return Deal(
        name="Test Acquisition",
        client_name="Test Corp",
        deal_type="pe_investment",
        deal_value="100000000",
        client_side="investor",
        jurisdiction="india",
        industry="fintech",
    )


def _make_advisory():
    return {
        "recommendation": "negotiate",
        "recommendation_rationale": "Several issues need resolution.",
        "executive_summary": "Test executive summary.",
        "financial_exposure": {
            "estimated_total": 5000000,
            "exposure_percentage": 5.0,
            "breakdown": [
                {"clause": "S.4.1", "risk_level": "high", "amount": 5000000},
            ],
        },
        "negotiation_priority_matrix": [],
        "negotiation_playbook": [
            {
                "issue": "[S.4.1] Low indemnity cap",
                "priority": "must_have",
                "opening_position": "Request 20% cap",
                "fallback_position": "Accept 15%",
                "walk_away_point": "Below 10%",
            },
        ],
        "key_risks": ["[S.4.1] Low indemnification cap"],
        "cross_clause_risks": ["[HIGH] Indemnification vs Reps: misalignment"],
        "timeline_risks": ["[S.3.1] Regulatory approval delay"],
        "recommendations": ["Increase indemnity cap to 15%"],
        "negotiation_points": [],
        "deal_breakers": [],
        "missing_clauses": [
            {"clause_type": "anti_dilution", "severity": "must_have", "rationale": "Required for PE investment"},
        ],
        "regulatory": {
            "checks": [
                {"title": "FEMA Compliance", "legal_basis": "FEMA NDI Rules 2019", "risk_level": "high"},
            ],
            "approval_sequence": [
                {"step": "Board Resolution", "timeline": "1-2 weeks", "depends_on": []},
                {"step": "FEMA Pricing", "timeline": "2-3 weeks", "depends_on": ["Board Resolution"]},
            ],
            "industry_risks": [
                {"regulator": "RBI", "description": "Digital lending compliance required"},
            ],
        },
        "risk_distribution": {"critical": 1, "high": 2, "medium": 3},
        "total_clauses_reviewed": 15,
        "human_override_rate": 0.1,
    }


def _make_contract_and_review():
    clause = Clause(
        id="cl1",
        contract_id="c1",
        clause_type=ClauseType.INDEMNIFICATION,
        title="Indemnification",
        section_reference="Section 4.1",
        text="The indemnification cap shall be 5% of the purchase price.",
    )
    contract = Contract(
        id="c1",
        deal_id="deal1",
        filename="test_sha.pdf",
        title="Share Holders Agreement",
        clauses=[clause],
    )
    finding = AIFinding(
        category="cap_analysis",
        title="Low indemnity cap",
        description="Cap at 5% is below market median of 15%.",
        risk_level=RiskLevel.HIGH,
        confidence=0.85,
        suggested_revision="The indemnification cap shall be 15% of the Aggregate Consideration.",
        market_comparison="Market median: 15% for Indian PE deals.",
        enforceability_note="",
    )
    cr = ClauseReview(
        clause_id="cl1",
        review_id="r1",
        stage=ReviewStage.APPROVED,
        ai_risk_level=RiskLevel.HIGH,
        ai_findings=[finding],
        ai_summary="Indemnity cap below market at 5%.",
        final_risk_level=RiskLevel.HIGH,
        human_annotations=[
            HumanAnnotation(
                reviewer_id="rev1",
                verdict=HumanVerdict.MODIFY,
                comment="Agree, push for 15%.",
                suggested_language="Cap at 15% of Aggregate Consideration.",
            ),
        ],
    )
    review = Review(
        id="r1",
        contract_id="c1",
        deal_id="deal1",
        clause_reviews=[cr],
    )
    return contract, review


class TestAdvisoryDocx:
    def test_generates_valid_docx(self):
        deal = _make_deal()
        advisory = _make_advisory()
        context = DealContext(client_side="investor", jurisdiction="india", industry="fintech")

        docx_bytes = export_advisory_docx(deal, advisory, [], context)
        assert isinstance(docx_bytes, bytes)
        assert len(docx_bytes) > 0

        # Verify it's a valid docx by parsing it
        doc = Document(BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Test Acquisition" in full_text
        assert "PROCEED WITH NEGOTIATIONS" in full_text

    def test_includes_playbook(self):
        deal = _make_deal()
        advisory = _make_advisory()
        docx_bytes = export_advisory_docx(deal, advisory, [])

        doc = Document(BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Negotiation Playbook" in full_text

    def test_includes_missing_clauses(self):
        deal = _make_deal()
        advisory = _make_advisory()
        docx_bytes = export_advisory_docx(deal, advisory, [])

        doc = Document(BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Missing Clauses" in full_text

    def test_includes_regulatory_roadmap(self):
        deal = _make_deal()
        advisory = _make_advisory()
        docx_bytes = export_advisory_docx(deal, advisory, [])

        doc = Document(BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Regulatory Approval Roadmap" in full_text


class TestRedlineDocx:
    def test_generates_valid_redline(self):
        contract, review = _make_contract_and_review()
        docx_bytes = export_redline_docx(contract, review, negotiation_round=1)

        assert isinstance(docx_bytes, bytes)
        assert len(docx_bytes) > 0

        doc = Document(BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Indemnification" in full_text
        assert "Round: 1" in full_text or "Negotiation Round: 1" in full_text

    def test_includes_findings(self):
        contract, review = _make_contract_and_review()
        docx_bytes = export_redline_docx(contract, review)

        doc = Document(BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Low indemnity cap" in full_text

    def test_includes_human_annotations(self):
        contract, review = _make_contract_and_review()
        docx_bytes = export_redline_docx(contract, review)

        doc = Document(BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "MODIFY" in full_text
        assert "15%" in full_text


class TestComplianceTracker:
    def test_generates_valid_tracker(self):
        deal = _make_deal()
        advisory = _make_advisory()
        docx_bytes = export_compliance_tracker_docx(deal, advisory)

        assert isinstance(docx_bytes, bytes)
        assert len(docx_bytes) > 0

        doc = Document(BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Compliance Tracker" in full_text
        assert "Test Acquisition" in full_text

    def test_includes_approvals(self):
        deal = _make_deal()
        advisory = _make_advisory()
        docx_bytes = export_compliance_tracker_docx(deal, advisory)

        doc = Document(BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Regulatory Approval" in full_text or "Pre-Closing" in full_text
